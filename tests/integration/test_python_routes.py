"""纯 Python 后端真实转换（PDF / 表格）。"""

from __future__ import annotations

from docmorph.formats import Format
from docmorph.results import ConflictPolicy


def convert(engine, source, target_format, tmp_path, name: str, **options):
    from docmorph.engine import make_request

    destination = tmp_path / f"{name}.{target_format}"
    request = make_request(
        source,
        destination,
        target_format,
        settings=engine.settings,
        conflict=ConflictPolicy.OVERWRITE,
        options=options,
    )
    return engine.convert(request)


def test_pdf_to_text_extracts_content(engine, samples, tmp_path):
    result = convert(engine, samples["pdf"], "txt", tmp_path, "pdf_out")
    assert result.ok, result.error
    assert "fixture" in result.primary_output.read_text(encoding="utf-8")


def test_pdf_to_docx_preserves_paragraphs(engine, samples, tmp_path):
    result = convert(engine, samples["pdf"], "docx", tmp_path, "pdf_docx")
    assert result.ok, result.error
    from docx import Document

    document = Document(str(result.primary_output))
    assert any("fixture" in paragraph.text for paragraph in document.paragraphs)


def test_pdf_to_pptx_renders_images_and_notes(engine, samples, tmp_path):
    result = convert(engine, samples["pdf"], "pptx", tmp_path, "pdf_pptx")
    assert result.ok, result.error
    from pptx import Presentation

    presentation = Presentation(str(result.primary_output))
    assert len(presentation.slides) == 1
    assert any("图片" in warning for warning in result.warnings)
    notes = presentation.slides[0].notes_slide.notes_text_frame.text
    assert "fixture" in notes


def test_xlsx_to_csv_splits_sheets(engine, samples, tmp_path):
    result = convert(engine, samples["xlsx"], "csv", tmp_path, "book")
    assert result.ok, result.error
    names = sorted(path.name for path in result.outputs)
    assert names == ["book.csv", "book__第二页.csv"]
    first = result.outputs[0].read_text(encoding="utf-8-sig")
    assert "苹果" in first
    assert "第二张表" not in first
    assert any("拆分" in warning for warning in result.warnings)


def test_xlsx_to_csv_merged_mode(engine, samples, tmp_path):
    result = convert(engine, samples["xlsx"], "csv", tmp_path, "merged", xlsx_csv_mode="merged")
    assert result.ok, result.error
    content = result.primary_output.read_text(encoding="utf-8-sig")
    assert "苹果" in content and "第二张表" in content
    assert "sheet" in content.splitlines()[0]


def test_xlsx_to_csv_first_only_mode_warns(engine, samples, tmp_path):
    result = convert(engine, samples["xlsx"], "csv", tmp_path, "first", xlsx_csv_mode="first")
    assert result.ok
    assert len(result.outputs) == 1
    assert any("仅导出" in warning for warning in result.warnings)


def test_xlsx_to_markdown_and_html_and_text(engine, samples, tmp_path):
    md = convert(engine, samples["xlsx"], "md", tmp_path, "book_md")
    assert md.ok, md.error
    md_text = md.primary_output.read_text(encoding="utf-8")
    assert md_text.startswith("## 数据")
    assert "| 名称 | 数量 | 单价 |" in md_text

    html = convert(engine, samples["xlsx"], "html", tmp_path, "book_html")
    assert html.ok
    assert "<table" in html.primary_output.read_text(encoding="utf-8")

    text = convert(engine, samples["xlsx"], "txt", tmp_path, "book_txt")
    assert text.ok
    assert "苹果" in text.primary_output.read_text(encoding="utf-8")


def test_csv_to_xlsx_creates_workbook(engine, samples, tmp_path):
    result = convert(engine, samples["csv"], "xlsx", tmp_path, "csv_out")
    assert result.ok, result.error
    from openpyxl import load_workbook

    workbook = load_workbook(str(result.primary_output))
    sheet = workbook.active
    assert sheet.cell(row=1, column=1).value == "名称"
    assert sheet.cell(row=2, column=1).value == "苹果"


def test_python_backend_handles_docx_to_text_standalone(engine, samples, tmp_path):
    """pandoc 缺失时的兜底路径：python 后端独立完成 docx→txt。"""
    from docmorph.backends.python_backend import PythonBackend
    from docmorph.settings.paths import TempWorkspace

    backend = PythonBackend()
    destination = tmp_path / "fallback.txt"
    with TempWorkspace(root=tmp_path, prefix="fallback") as workspace:
        outcome = backend.convert(
            samples["docx"], destination, Format.DOCX, Format.TXT, {}, workspace
        )
    assert outcome.outputs == [destination]
    content = destination.read_text(encoding="utf-8")
    assert "中文正文" in content
    assert "单元格A1" in content
