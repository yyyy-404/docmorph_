"""pandoc 后端真实转换（缺 pandoc 时自动跳过）。"""

from __future__ import annotations

from docmorph.formats import Format
from docmorph.registry import select_plan
from docmorph.results import ConflictPolicy
from tests.conftest import PANDOC_SKIP

pytestmark = PANDOC_SKIP


def convert(engine, source, target_format, tmp_path, name: str):
    from docmorph.engine import make_request

    destination = tmp_path / f"{name}.{target_format}"
    request = make_request(
        source,
        destination,
        target_format,
        settings=engine.settings,
        conflict=ConflictPolicy.OVERWRITE,
    )
    return engine.convert(request)


def test_docx_to_markdown_keeps_text_and_extracts_media(engine, samples, tmp_path):
    result = convert(engine, samples["docx"], "md", tmp_path, "docx_out")
    assert result.ok, result.error
    text = result.primary_output.read_text(encoding="utf-8")
    assert "DocMorph 测试文档" in text
    assert "中文正文" in text
    assert "单元格A1" in text
    assert any("图片" in warning for warning in result.warnings)
    assert (tmp_path / "docx_out_media").is_dir()


def test_docx_to_html_and_txt(engine, samples, tmp_path):
    html = convert(engine, samples["docx"], "html", tmp_path, "docx_html")
    assert html.ok, html.error
    assert "中文正文" in html.primary_output.read_text(encoding="utf-8")

    text = convert(engine, samples["docx"], "txt", tmp_path, "docx_txt")
    assert text.ok, text.error
    content = text.primary_output.read_text(encoding="utf-8")
    assert "中文正文" in content
    # 回归：过去用 pandoc 的 "txt" 输出格式必然失败，现映射为 plain
    assert "中文正文" in content


def test_markdown_roundtrip_to_docx_and_html(engine, samples, tmp_path):
    docx = convert(engine, samples["md"], "docx", tmp_path, "md_docx")
    assert docx.ok, docx.error
    from docx import Document

    document = Document(str(docx.primary_output))
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Markdown 标题" in all_text
    assert document.tables, "Markdown 表格应转换为 docx 表格"

    html = convert(engine, samples["md"], "html", tmp_path, "md_html")
    assert html.ok, html.error
    assert "Markdown 标题" in html.primary_output.read_text(encoding="utf-8")


def test_html_to_docx_and_markdown(engine, samples, tmp_path):
    docx = convert(engine, samples["html"], "docx", tmp_path, "html_docx")
    assert docx.ok, docx.error
    md = convert(engine, samples["html"], "md", tmp_path, "html_md")
    assert md.ok, md.error
    assert "HTML 标题" in md.primary_output.read_text(encoding="utf-8")


def test_text_file_uses_markdown_reader(engine, samples, tmp_path):
    docx = convert(engine, samples["txt"], "docx", tmp_path, "txt_docx")
    assert docx.ok, docx.error
    assert docx.primary_output.stat().st_size > 0


def test_csv_to_docx_and_html(engine, samples, tmp_path):
    docx = convert(engine, samples["csv"], "docx", tmp_path, "csv_docx")
    assert docx.ok, docx.error
    html = convert(engine, samples["csv"], "html", tmp_path, "csv_html")
    assert html.ok, html.error
    assert "苹果" in html.primary_output.read_text(encoding="utf-8")


def test_xlsx_to_docx_chain(engine, samples, tmp_path):
    """xlsx → md（python）→ docx（pandoc）：验证链式方案。"""
    result = convert(engine, samples["xlsx"], "docx", tmp_path, "xlsx_docx")
    assert result.ok, result.error
    assert result.backend == "python + pandoc"
    from docx import Document

    document = Document(str(result.primary_output))
    # 表格内容会转成 docx 表格（而不是普通段落）
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "苹果" in cells


def test_every_declared_target_has_available_plan(engine):
    for source in (Format.DOCX, Format.MD, Format.HTML, Format.TXT, Format.CSV):
        for target in engine.declared_targets(source):
            plan, reason = select_plan(source, target, engine.backends, engine.capabilities)
            assert plan is not None, f"{source}->{target} 无可用方案：{reason}"
