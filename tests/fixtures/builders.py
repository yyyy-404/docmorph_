"""测试文档生成器：用代码造样本，避免在仓库里提交二进制文件。"""

from __future__ import annotations

from pathlib import Path

CJK_TEXT = "这是一段中文正文，用于验证转换后的文本完整性。"
TABLE_MARK = "单元格A1"


def make_docx(path: Path, with_image: bool = True, with_table: bool = True) -> Path:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_heading("DocMorph 测试文档", level=1)
    document.add_paragraph(CJK_TEXT)
    document.add_heading("结构与内容", level=2)
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = TABLE_MARK
        table.cell(0, 1).text = "单元格B1"
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "2"
    if with_image:
        from PIL import Image

        image_path = path.parent / f"{path.stem}-pixel.png"
        Image.new("RGB", (120, 60), (140, 192, 235)).save(image_path)
        document.add_picture(str(image_path), width=Inches(1.0))
    document.save(str(path))
    return path


def make_xlsx(path: Path, sheets: int = 2) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    first = workbook.active
    first.title = "数据"
    first.append(["名称", "数量", "单价"])
    first.append(["苹果", 3, 4.5])
    first.append(["香蕉", 5, 2.0])
    first.append(["合计", "=SUM(B2:B3)", "=SUMPRODUCT(B2:B3,C2:C3)"])
    if sheets > 1:
        second = workbook.create_sheet("第二页")
        second.append(["第二张表", "内容"])
    workbook.save(str(path))
    return path


def make_markdown(path: Path) -> Path:
    path.write_text(
        "# Markdown 标题\n\n"
        f"{CJK_TEXT}\n\n"
        "- 项目一\n- 项目二\n\n"
        "| 列1 | 列2 |\n|---|---|\n| a | b |\n",
        encoding="utf-8",
    )
    return path


def make_html(path: Path) -> Path:
    path.write_text(
        "<!doctype html><html><head><meta charset='utf-8'><title>DocMorph</title></head>"
        f"<body><h1>HTML 标题</h1><p>{CJK_TEXT}</p>"
        "<table><tr><td>a</td><td>b</td></tr></table></body></html>",
        encoding="utf-8",
    )
    return path


def make_text(path: Path, text: str = "纯文本第一行\n第二行 Hello DocMorph\n") -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def make_csv(path: Path) -> Path:
    path.write_text("名称,数量\n苹果,3\n香蕉,5\n", encoding="utf-8-sig")
    return path


def make_pdf(path: Path, text: str = "DocMorph PDF fixture line one") -> Path:
    """用 PyMuPDF 生成一个可抽取文本的 PDF（不依赖 Office）。"""
    from docmorph.backends.python_backend import import_pymupdf

    pymupdf = import_pymupdf()
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 96), text, fontsize=14)
    page.insert_text((72, 130), "second line of the fixture document", fontsize=12)
    document.save(str(path))
    document.close()
    return path


def build_all(directory: Path) -> dict[str, Path]:
    """生成全套样本文件，返回 ``{格式: 路径}``。"""
    directory.mkdir(parents=True, exist_ok=True)
    return {
        "docx": make_docx(directory / "sample.docx"),
        "docx_plain": make_docx(directory / "sample_plain.docx", with_image=False, with_table=False),
        "xlsx": make_xlsx(directory / "sample.xlsx"),
        "md": make_markdown(directory / "sample.md"),
        "html": make_html(directory / "sample.html"),
        "txt": make_text(directory / "sample.txt"),
        "csv": make_csv(directory / "sample.csv"),
        "pdf": make_pdf(directory / "sample.pdf"),
    }
